"""README_DEVELOPER.md, OPS_NOTES.md and decoy_summary.xlsx for one locale's alignment decoy package."""
import collections, random

LANG_NAME = {"de_DE": "German", "es_ES": "Spanish (Spain)", "it_IT": "Italian", "pt_BR": "Portuguese (Brazil)", "fr_FR": "French"}
BAND_TEXT = {"A1": "Obvious - missing it costs the most", "A2": "Medium", "A3": "Subtle - costs the least"}
RULES = [  # family, band rule, example
    ("BND-MINOR", "A3: edge moved 40-80 ms", "start 60 ms late, inside the word"),
    ("BND-MAJOR", "A2: edge moved 100-300 ms · A1: more than 300 ms", "start 500 ms early, in the silence before the word"),
    ("BND-OVERLAP", "A1: edge runs into the next / previous word (same-speaker overlap)", "end 150 ms inside the next word"),
    ("BND-ZERO", "A1: start = end", "segment with 0 ms length"),
    ("WD", "A1: content word · A2: function word / short word", "spoken word with no segment"),
    ("WH", "A1: segment over silence · A2: segment over a breath / sigh", "invented word over a pause"),
    ("WS", "A1: different word · A2: homophone from the guideline · A3: small form change", "seite -> saite, hola -> ola"),
    ("SEG-MERGE", "A2", "two words in one segment"),
    ("SEG-SPLIT", "A2", "compound or elided form cut into two segments"),
    ("SEG-LETTERS", "A2", "spelled letters joined: a r d -> ard"),
    ("SEG-FILLER", "A2", "filler hidden inside the next word's segment"),
    ("G-NUM", "A1", "number written as digits"),
    ("G-CASE", "A2", "capital letter"),
    ("G-HYPHEN", "A2", "hyphenated word kept as one segment"),
    ("G-ACCENT", "A2 umlaut written ae/oe/ue/ss · A3 accent dropped", "für -> fuer, está -> esta"),
    ("G-CONTRACT", "A2", "al -> a el, do -> de o"),
    ("G-BACKCH", "A3", "filler misspelt, or mm hm joined"),
    ("G-PUNCT", "A2", "comma left in a segment"),
    ("CB-WHOLE", "A1 · fixed 5 points, outside the 100", "entirely synthetic file not marked"),
]


def counts(corpus):
    fam = collections.Counter(x["family"] for e in corpus for x in e["injections"])
    band = collections.Counter(x["band"] for e in corpus for x in e["injections"])
    return fam, band


def write_xlsx(path, loc, corpus, FAMILY, GROUP_POINTS, BAND_FACTOR, CB_WEIGHT):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="1F4E78"); wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head, fill, wrap
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        ws.freeze_panes = "A2"

    fam, band = counts(corpus)
    ws = wb.active; ws.title = "Summary"
    n_inj = sum(fam.values())
    syn = sum(e["voice"]["type"] == "synthetic" for e in corpus)
    ws.append([f"Alignment decoy corpus · {LANG_NAME[loc]} ({loc})"]); ws["A1"].font = Font(bold=True, size=14)
    for row in ([], ["Files", len(corpus)], ["Files with decoys", len(corpus)], ["Synthetic files (answer key only)", syn],
                ["Human-sounding files", len(corpus) - syn], ["Injections", n_inj],
                ["Alignment points (sum of weights, excl. CB-WHOLE)", round(sum(x["weight"] for e in corpus for x in e["injections"] if x["family"] != "CB-WHOLE"), 1)],
                [], ["Score group", "Share", "Injections", "Points"]):
        ws.append(row)
    ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = ws.cell(ws.max_row, 3).font = ws.cell(ws.max_row, 4).font = Font(bold=True)
    for g, pts in GROUP_POINTS.items():
        k = sum(1 for e in corpus for x in e["injections"] if FAMILY[x["family"]][0] == g)
        p = sum(x["weight"] for e in corpus for x in e["injections"] if FAMILY[x["family"]][0] == g)
        ws.append([g, f"{pts}%", k, round(p, 1)])
    ws.append(["CHECKBOX (CB-WHOLE)", "gate", fam.get("CB-WHOLE", 0), round(CB_WEIGHT * fam.get("CB-WHOLE", 0), 1)])
    ws.append([]); ws.append(["Band", "Meaning", "Injections", "Weight factor"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    for b in ("A1", "A2", "A3"):
        ws.append([b, BAND_TEXT[b], band.get(b, 0), BAND_FACTOR[b]])
    ws.column_dimensions["A"].width = 48; ws.column_dimensions["B"].width = 40; ws.column_dimensions["C"].width = 12; ws.column_dimensions["D"].width = 14

    weights = collections.defaultdict(set)
    for e in corpus:
        for x in e["injections"]:
            weights[(x["family"], x["band"])].add(x["weight"])
    rows = []
    for f, rule, ex in RULES:
        g, det, name = FAMILY[f]
        wb_ = "; ".join(f"{b}: {sorted(weights[(f, b)])[0]}" for b in ("A1", "A2", "A3") if weights[(f, b)])
        rows.append([f, name, g, det, rule, wb_ or "-", fam.get(f, 0), ex])
    sheet(wb.create_sheet("Bands & weights"), ["Family", "What is planted", "Score group", "detect", "Band rule",
                                                "Weight per injection", "Count", "Example"], rows, [13, 42, 15, 27, 44, 22, 8, 36])

    rows = [[e["decoy_id"], e["file_name"], e["file_id"], e["voice"]["type"], e["gold_whole_audio_tag"], e["voice"]["name"],
             e["duration_seconds"], e["injection_count"], ", ".join(e["families"]), ", ".join(e["bands"]), e["total_weight"],
             e["gold_text"], e["decoy_text"]] for e in corpus]
    sheet(wb.create_sheet("Files"), ["decoy_id", "file_name", "file_id", "voice_type", "gold_whole_audio_tag", "voice", "duration_s",
                                     "injections", "families", "bands", "total_weight", "gold_text", "decoy_text"],
          rows, [14, 46, 28, 11, 18, 30, 10, 10, 34, 10, 11, 60, 60])

    rows = []
    for e in corpus:
        for x in e["injections"]:
            m = x.get("meta", {})
            g, d = m.get("gold") or {}, m.get("decoy") or {}
            what = (m.get("word") or (f"{m.get('from')} -> {m.get('to')}" if "from" in m else "")
                    or (m.get("joined") or "") or f"{x.get('gold')} -> {x.get('decoy')}")
            rows.append([x["injection_id"], e["file_name"], x["family"], FAMILY[x["family"]][0], x["band"], x["weight"],
                         m.get("rule") or x.get("rule"), what, g.get("start"), g.get("end"), d.get("start"), d.get("end"),
                         m.get("shift_ms"), m.get("left", ""), m.get("right", ""), m.get("plain_offset", ""), x["fingerprint"]])
    sheet(wb.create_sheet("Injections"), ["injection_id", "file_name", "family", "score_group", "band", "weight", "rule", "target",
                                          "gold_start", "gold_end", "decoy_start", "decoy_end", "shift_ms", "left_anchor",
                                          "right_anchor", "plain_offset", "fingerprint"],
          rows, [18, 46, 13, 15, 6, 8, 26, 22, 10, 10, 10, 10, 9, 34, 34, 11, 18])

    rng = random.Random(f"spot-{loc}")
    pick = rng.sample(rows, min(15, len(rows)))
    sheet(wb.create_sheet("Spot-check"), ["injection_id", "file_name", "family", "band", "rule", "target", "gold_start", "gold_end",
                                          "decoy_start", "decoy_end", "checked_by", "ok?", "notes"],
          [[r[0], r[1], r[2], r[4], r[6], r[7], r[8], r[9], r[10], r[11], "", "", ""] for r in pick],
          [18, 46, 13, 6, 26, 22, 10, 10, 10, 10, 14, 8, 30])
    wb.save(path)


def write_docs(pkg, loc, corpus, st, scorer_md5, FAMILY, GROUP_POINTS, CB_WEIGHT, SEED):
    fam, band = counts(corpus)
    lang = LANG_NAME[loc]
    syn = sum(e["voice"]["type"] == "synthetic" for e in corpus)
    n_inj = sum(fam.values())
    fam_rows = "\n".join(f"| `{f}` | {FAMILY[f][2]} | {FAMILY[f][0]} | `{FAMILY[f][1]}` | {rule} | {fam.get(f, 0)} |"
                         for f, rule, _ in RULES)
    group_rows = "\n".join(f"| {g} | {p}% | {sum(1 for e in corpus for x in e['injections'] if FAMILY[x['family']][0] == g)} |"
                           for g, p in GROUP_POINTS.items())
    e0 = next(e for e in corpus if any(x["family"].startswith("BND") for x in e["injections"]))
    b0 = next(x for x in e0["injections"] if x["family"].startswith("BND"))

    readme = f"""# Alignment Decoy Programme — Developer README

**v1 · first build · {lang} (`{loc}`) · 20 files · {n_inj} injections · seed `{SEED}`**

This is the forced-alignment counterpart of the transcription decoy package. Same idea, same file set:
we know exactly what we changed in the machine pre-segmentation, so catching it is exact and evidenced.

---

## Files in this handover

| File | What to do with it |
|---|---|
| `alignment_prefill_{loc}.json` | **Load this into the alignment tool as the machine pre-segmentation** of these 20 files. Sandbox format: `fileMetadata` + `items[]` (`id`, `name`, `transcript`, `source.words[]` with `word` / `start` / `end` in seconds). The planted errors are in here — word text and word start/end. Nothing in it marks a file as a decoy or as synthetic. |
| `audio/*.wav` | The 20 audio files (16 kHz mono). File names match `items[].name`. |
| `decoy_corpus.json` | **The answer key. Server-side only.** One record per file: gold vs decoy text and segments, whole-audio tag, every injection with family, band, weight, anchors and fingerprint. |
| `alignmentDecoyScorer.js` | The scoring engine. Zero dependencies. **Do not edit.** MD5 `{scorer_md5}` |
| `run_selftest.js` | `node run_selftest.js decoy_corpus.json` — re-runs the verification below. Output of this build: `selftest_report.txt`. |
| `decoy_corpus_ALL.csv` | Reference — all {len(corpus)} files, {n_inj} rows, one per injection |
| `decoy_corpus_CHANGED.csv` | Reference — files changed since the previous version. **Empty in v1** (header only) |
| `decoy_change_manifest.json` | The four ID lists: fixed, reinjected, pulled, flagged. **All empty in v1** |
| `decoy_summary.xlsx` | Bands & weights, score groups, every file, every injection, a 15-row spot-check sheet |
| `OPS_NOTES.md` | Full integration guide, output schemas, verdicts, bands, rotation policy |
| `Forced_Alignment_Guidelines_{loc}.pdf` | The annotator guideline the decoys are built on |

---

## Do this

- [ ] **Load `alignment_prefill_{loc}.json` as the pre-segmentation** for these 20 files, exactly as given.
      The timings are the test: if the tool re-runs its own aligner on these files, every boundary decoy disappears.
- [ ] **Load the corpus server-side and filter on `status`:**
  ```js
  const S = require('./alignmentDecoyScorer.js');
  const corpus = JSON.parse(fs.readFileSync('./decoy_corpus.json','utf8')).filter(c => c.status === 'ACTIVE');
  const decoys = S.indexCorpus(corpus);              // lookup by file_id (= items[].id) or decoy_id
  ```
- [ ] **On submit**, score it (see the submission shape below) and store the result:
  ```js
  if (decoys.isDecoy(fileId)) {{
    const result = S.scoreDecoy(submission, decoys.get(fileId), {{ annotator_id, submitted_at, batch_id, seconds_spent }});
    await db.alignment_decoy_results.insert(result);
  }}
  ```
- [ ] **Serve the 20 files mixed into normal work**, in random order — never as a block.
- [ ] **Keep `decoy_corpus.json` off every front-end route** — it holds the gold alignment and says which files are synthetic.

## Do NOT do this

- Do not show the voice, the `group` (`AL-HUMAN` / `AL-SYNTH`) or anything from `decoy_corpus.json` in the tool
- Do not edit `alignmentDecoyScorer.js` — verified against this corpus (below)
- Do not let the tool snap or re-align the pre-filled timings of these files

---

## Submission shape

`scoreDecoy` accepts the submitted segments in any of these shapes:

```js
[{{ word: 'zug', start: 1.23, end: 1.51 }}, ...]                                 // segments only
{{ segments: [...], whole_audio_tag: 'none' | 'synthetic_speech', discard_reason: null }}
{{ transcript: '...', source: {{ words: [...] }}, whole_audio_tag: 'none' }}      // an item in the pre-fill format
```

`whole_audio_tag` is the dustbin answer "Entirely synthetic speech" (guideline 2.5). Pass `'synthetic_speech'` when the
annotator marked the file as entirely synthetic, `'none'` otherwise. Leave it out and the synthetic check is
`INDETERMINATE` (never a miss). `discard_reason` is any other dustbin reason.

---

## Fields on each corpus record

| Field | Meaning |
|---|---|
| `decoy_id` | `ALN-{loc[:2].upper()}-0001` … — pre-shuffled, not in serve order |
| `file_id`, `file_name` | Same as `items[].id` / `items[].name` in the pre-fill and the wav name |
| `gold_text`, `decoy_text` | Correct transcript vs what the pre-fill shows (space-separated segments) |
| `gold_segments`, `decoy_segments` | Correct alignment vs the pre-fill alignment (`word`, `start`, `end`) |
| `gold_whole_audio_tag` | `synthetic_speech` for the {syn} entirely synthetic files, else `none`. The pre-fill never says it (`decoy_whole_audio_tag` = `none`) |
| `voice` | Voice name, gender and `type` (`human` / `synthetic`) — answer key only |
| `injections[]` | `family`, `band`, `weight`, `meta` (`detect`, `rule`, `word` / `from`→`to`, gold and decoy span, `shift_ms`, `left` / `right` anchors, `plain_offset`), `injection_id`, `fingerprint` |
| `status`, `cat1b_flag`, `cat2_fixes`, `reinjected` | As in the transcription corpus. v1: all `ACTIVE`, no flags, no fixes |

---

## Current state

| | |
|---|---|
| Files | **{len(corpus)}** ({len(corpus) - syn} human-sounding, {syn} entirely synthetic) — every file carries decoys |
| Injections | **{n_inj}** ({st['alignment']} alignment + {st['checkbox']} synthetic check) · A1 {band.get('A1', 0)} · A2 {band.get('A2', 0)} · A3 {band.get('A3', 0)} |
| Verification | {st['decoy_missed']}/{st['scored']} `MISSED` on decoy · {st['gold_caught']}/{st['scored']} `CAUGHT` on gold · {st['jitter_caught']}/{st['alignment']} `CAUGHT` on gold with ±20 ms jitter |
| Throughput | {st['ms_per_submission']} ms per submission — call it inline, no queue |
"""

    ops = f"""# Alignment Decoy Programme — Integration & Ops

**Corpus:** {len(corpus)} files · {n_inj} injections · {lang} (`{loc}`) · seed `{SEED}`
**Source:** 20 new transcripts written to the `{loc}` forced-alignment guideline, voiced with ElevenLabs
(17 human-sounding voices with breathing and pauses, {syn} entirely synthetic), gold alignment by forced alignment
plus loudness-based correction at pauses, checked by speech-to-text and a boundary audit.

---

## 1. Files

| File | What it is |
|---|---|
| `alignment_prefill_{loc}.json` | What the annotator is served: the pre-segmentation with the planted errors. |
| `audio/` | The 20 wav files. |
| `decoy_corpus.json` | The corpus and answer key. **Never expose it to the annotation front end.** |
| `decoy_corpus_ALL.csv` | Same data, one row per injection, for spreadsheets and SQL. |
| `alignmentDecoyScorer.js` | The scoring engine. Zero dependencies. |
| `decoy_summary.xlsx` | Bands & weights, score groups, files, injections, spot-check sample. |

---

## 2. Wiring it up

```js
const fs = require('fs');
const S  = require('./alignmentDecoyScorer.js');
const corpus = JSON.parse(fs.readFileSync('./decoy_corpus.json', 'utf8')).filter(c => c.status === 'ACTIVE');
const decoys = S.indexCorpus(corpus);
```

### When a task is served

```js
if (decoys.isDecoy(fileId)) {{
  servePreSegmentation(prefill.items.find(i => i.id === fileId));   // the planted pre-segmentation, as is
}} else {{
  servePreSegmentation(normalAligner(fileId));
}}
```

### When a submission arrives

```js
if (decoys.isDecoy(fileId)) {{
  const entry  = decoys.get(fileId);
  const result = S.scoreDecoy({{ segments, whole_audio_tag, discard_reason }}, entry, {{
    annotator_id: userId, submitted_at: new Date().toISOString(), batch_id: batchId, seconds_spent: secondsSpent,
  }});
  await db.alignment_decoy_results.insert(result);
  await db.alignment_decoy_residuals.insert(S.classifyResiduals({{ segments }}, entry));   // advisory only
}}
```

### The daily pull

```js
const report = S.buildDailyReport(rows, {{ from, to }});        // all-time within range
const recent = S.buildDailyReport(rows, {{ window: 20 }});      // last 20 decoys per person
fs.writeFileSync('daily.csv', S.reportToCsv(report));
const detail = S.injectionsToRows(rows);                      // per-injection audit table
```

---

## 3. Output shapes

### `scoreDecoy` — per submission

```json
{{
  "decoy_id": "{e0['decoy_id']}", "group": "{e0['group']}", "file_id": "{e0['file_id']}",
  "annotator_id": "annotator7230", "injection_count": {e0['injection_count']}, "caught": {e0['injection_count'] - 1}, "missed": 1, "indeterminate": 0,
  "excluded": 0, "file_is_synthetic": {str(e0['voice']['type'] == 'synthetic').lower()}, "marked_synthetic": false, "false_synthetic": false,
  "points_possible": {e0['total_weight']}, "points_lost": {b0['weight']}, "score_percent": {round(100 * (1 - b0['weight'] / e0['total_weight']), 1)},
  "missed_families": ["{b0['family']}"], "missed_bands": ["{b0['band']}"],
  "by_band": {{ "A1": {{ ... }}, "A2": {{ ... }}, "A3": {{ ... }} }},
  "by_family": {{ ... }}, "by_group": {{ "BOUNDARY": {{ ... }}, "WRONG-WORD": {{ ... }} }},
  "injections": [ /* per-injection detail, below */ ]
}}
```

### Per-injection detail (a boundary decoy)

```json
{{
  "injection_id": "{b0['injection_id']}", "family": "{b0['family']}", "group": "BOUNDARY", "band": "{b0['band']}", "weight": {b0['weight']},
  "detect_rule": "{b0['meta']['detect']}", "verdict": "MISSED", "confidence": 1,
  "evidence": {{ "word": "{b0['meta']['word']}", "edge": "{b0['meta']['edge']}", "gold_s": {b0['meta']['gold'][b0['meta']['edge']]},
    "planted_s": {b0['meta']['decoy'][b0['meta']['edge']]}, "submitted_s": {b0['meta']['decoy'][b0['meta']['edge']]},
    "remaining_error_ms": {b0['meta']['shift_ms']}, "planted_shift_ms": {b0['meta']['shift_ms']} }}
}}
```

---

## 4. Verdicts

| Verdict | Meaning |
|---|---|
| `CAUGHT` | The annotator found and fixed the planted error |
| `MISSED` | The planted error survived into the submission |
| `INDETERMINATE` | Can't be judged (region rewritten, a third value, a partial fix). **Route to human review — never count as a miss.** |
| `EXCLUDED` | The file is entirely synthetic and the annotator correctly marked it so: per guideline 2.5 they leave it empty, so its alignment decoys are not scored |

**Timing rule.** An edge is `CAUGHT` when it is back within **30 ms** of gold. It is `MISSED` when at least half of the
planted shift is still there, in the planted direction. Anything in between is `INDETERMINATE`.

**Synthetic files.** Marked synthetic → `CB-WHOLE` `CAUGHT`, the rest `EXCLUDED`. Not marked → `CB-WHOLE` `MISSED` and
the alignment decoys are scored normally. A **human** file marked synthetic or discarded → every injection `MISSED`
and `false_synthetic: true` (the daily report counts these).

---

## 5. Bands, score groups and weights

Bands follow the transcription corpus: **A1** obvious (missing it costs the most), **A2** medium, **A3** subtle.
Boundary bands: minor 40–80 ms = **A3**; major 100–300 ms = **A2**; more than 300 ms, or running into the next word = **A1**.
Zero length = **A1**.

Each score group gets a fixed share of **100 alignment points per locale**, spread over its injections in the ratio
A1 3 : A2 2 : A3 1 (so an A1 miss costs three times an A3 miss). The synthetic check is a separate gate worth
{CB_WEIGHT:g} points per synthetic file. Planted counts follow the same mix.

| Score group | Share | Injections |
|---|---|---|
{group_rows}

| Family | What is planted | Group | `detect` | Band rule | Count |
|---|---|---|---|---|---|
{fam_rows}

---

## 6. Category A vs B

**Category A** is what this corpus measures: we know exactly what we changed. **Category B** (`classifyResiduals`) is a
diff against gold outside the injected regions — a coaching signal only (`advisory: true`). Gold timings are accurate
to about ±20–30 ms, so B ignores edge differences under 100 ms.

---

## 7. Anti-detection

- **Never expose `decoy_corpus.json` to the front end.** It is the answer key.
- **Serve decoys through the normal queue.** File names and ids are random (`<uuid>_NNN.wav`, `cm…`), the pre-fill lists
  files in an order unrelated to `decoy_id` or file type, and nothing in the pre-fill marks a decoy or a synthetic file.
- **The injections look like ordinary machine pre-segmentation errors**: edges a little off, a segment over a breath, a
  missing filler, digits, a joined acronym. Exactly the errors the guideline (2.3) tells annotators to fix.
- **Rotate.** Retire a decoy after roughly 30 exposures; regenerate with a new seed.

---

## 8. Known limits

- **Gold timing precision is about ±20–30 ms** (forced alignment plus loudness-based correction at pauses). That is why
  minor decoys start at 40 ms and the catch tolerance is 30 ms. The self-test shows gold with ±20 ms random jitter on
  every edge still scores `CAUGHT` {st['jitter_caught']}/{st['alignment']}.
- **The pre-segmentation must be served as given.** If the alignment tool re-aligns on load, boundary decoys
  (50 % of the points) vanish; transcript-level decoys (text, missing, hallucinated) still work.
- **Band assignment is a design choice, not a measurement.** Tune it against real miss rates after the first two weeks.
- **All audio is generated speech.** "Human-sounding" files use natural voices with breathing, sighs and uneven pauses;
  "synthetic" files are deliberately robotic (no breathing, exactly 0.16 s between words).

---

## 9. Reproducing the corpus

From the repository root:

```bash
python make_japanese.py --lang decoy-{loc}            # audio + gold alignment (ElevenLabs; skips clips already made)
python alignment_decoys/build_alignment_decoys.py {loc}   # plan + inject + package + self-test + zip
```

Seed `{SEED}` regenerates this corpus byte-identically from the same audio. Change the seed for a fresh set.
"""
    (pkg / "README_DEVELOPER.md").write_text(readme, encoding="utf-8")
    (pkg / "OPS_NOTES.md").write_text(ops, encoding="utf-8")
